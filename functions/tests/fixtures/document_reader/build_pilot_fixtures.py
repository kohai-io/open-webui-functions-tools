"""Rebuild the small synthetic upload fixtures with existing docx/reportlab packages."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate


HERE = Path(__file__).parent
TITLE = "Document reader pilot brief"
SECTIONS = [
    ("Purpose", "This synthetic brief tests document reading and source inspection. It proposes a four-week pilot with two teams. It does not record an approved project or a commitment to spend."),
    ("Scope", "The two teams will use synthetic documents only. Customer records, employee information and unpublished business material are outside the pilot scope."),
    ("Conditions", "The proposed budget is GBP 12,000, subject to approval by the project sponsor. No purchase may be made until that approval is recorded."),
    ("Timing", "The proposed start date is 2 November 2026. This date is provisional and depends on the access review being completed first."),
    ("Evaluation", "The team will aim to review 20 documents during the pilot. This is a target, not a promised outcome. The pilot may stop early if material source errors are found."),
    ("Next step", "The coordinator should record the access-review outcome and seek approval. The coordinator must not describe the pilot as approved before those conditions are met."),
]


def build():
    markdown = "# " + TITLE + "\n\n" + "\n\n".join("## " + title + "\n\n" + prose for title, prose in SECTIONS) + "\n"
    (HERE / "pilot-brief.md").write_text(markdown, encoding="utf-8", newline="\n")

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in (("Normal", 11), ("Title", 18), ("Heading 1", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    document.styles["Normal"].paragraph_format.line_spacing = 1.12
    document.styles["Heading 1"].paragraph_format.space_before = Pt(6)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    document.add_paragraph(TITLE, style="Title")
    for heading, prose in SECTIONS:
        document.add_paragraph(heading, style="Heading 1")
        document.add_paragraph(prose)
    document.core_properties.title = TITLE
    document.core_properties.subject = "Synthetic Document Reader upload fixture"
    document.core_properties.author = ""
    document.save(HERE / "pilot-brief.docx")

    styles = {
        "title": ParagraphStyle("pilot-title", fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=12, textColor=colors.black),
        "heading": ParagraphStyle("pilot-heading", fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=6, spaceAfter=4, textColor=colors.black),
        "body": ParagraphStyle("pilot-body", fontName="Helvetica", fontSize=11, leading=14, spaceAfter=8, textColor=colors.black),
    }
    story = [Paragraph(TITLE, styles["title"])]
    for heading, prose in SECTIONS:
        story.extend([Paragraph(heading, styles["heading"]), Paragraph(prose, styles["body"])])
    SimpleDocTemplate(str(HERE / "pilot-brief.pdf"), pagesize=letter, topMargin=54, bottomMargin=54, leftMargin=61.2, rightMargin=61.2, title=TITLE, author="", invariant=1).build(story)


if __name__ == "__main__":
    build()
